"""Regression tests for ``MLPlanDocHelper.extract_relevant_text`` decryption.

Before this fix, ``extract_relevant_text`` read ``plan_document_enc`` via
``file_field.path`` and handed the raw *ciphertext* file to the PDF parser, so
extraction always failed, the exception was swallowed, and the feature silently
yielded "".  These tests save a real encrypted PDF to ``plan_document_enc`` and
assert the decrypted text is returned instead.

These live in ``tests/sync/`` rather than ``tests/async/`` on purpose: they are
synchronous ``TransactionTestCase`` methods that drive the async helper via
``asyncio.run()``. ``TransactionTestCase`` truncates tables and cannot run under
the async suite's parallel (``-n auto``) xdist workers, and the sync suite is
where the ``asyncio.run()``-from-a-sync-test pattern belongs.
"""

import asyncio
from unittest import mock

import pymupdf
import pytest
from django.core.files.base import ContentFile
from django.test import TransactionTestCase

from fighthealthinsurance.ml.ml_plan_doc_helper import MLPlanDocHelper
from fighthealthinsurance.models import Denial, PlanDocuments


def _make_pdf_bytes(text: str) -> bytes:
    """Build a one-page PDF containing *text* and return its raw bytes."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    pdf_bytes: bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.mark.django_db
class ExtractRelevantTextDecryptionTests(TransactionTestCase):
    def _make_denial(self) -> Denial:
        return Denial.objects.create(
            hashed_email="hash:plan-doc",
            denial_text="Sample denial for plan document extraction.",
        )

    def test_extract_relevant_text_decrypts_encrypted_pdf(self):
        """An encrypted PDF whose text matches a term yields decrypted text."""
        denial = self._make_denial()
        pdf_bytes = _make_pdf_bytes(
            "This plan covers services subject to medical necessity review."
        )
        doc = PlanDocuments(denial=denial)
        # EncryptedFileField encrypts the bytes on save.
        doc.plan_document_enc.save("plan.pdf", ContentFile(pdf_bytes), save=True)

        result = asyncio.run(
            MLPlanDocHelper.extract_relevant_text(
                denial.denial_id, ["medical necessity"]
            )
        )

        # Prior to the decryption fix this returned "" because the ciphertext
        # was handed to the PDF parser. The decrypted text must now come back
        # (assertIn also proves the result is non-empty).
        self.assertIn("medical necessity", result.lower())

    def test_extract_relevant_text_skips_docs_without_matching_terms(self):
        """A decrypted doc with no matching term contributes no text."""
        denial = self._make_denial()
        pdf_bytes = _make_pdf_bytes("This document is only about dental cleanings.")
        doc = PlanDocuments(denial=denial)
        doc.plan_document_enc.save("plan.pdf", ContentFile(pdf_bytes), save=True)

        result = asyncio.run(
            MLPlanDocHelper.extract_relevant_text(
                denial.denial_id, ["prior authorization"]
            )
        )

        self.assertEqual(result, "")


class WhatItWillAndWillNotParseTest(TransactionTestCase):
    """The parser runs on a shared worker, so what reaches it is bounded."""

    def test_a_document_over_the_ceiling_is_skipped_not_parsed(self):
        from unittest.mock import patch

        denial = Denial.objects.create(
            hashed_email=Denial.get_hashed_email("big@example.com"),
            denial_text="Denied.",
        )
        oversized = b"%PDF-1.4\n" + b"x" * (MLPlanDocHelper.MAX_DOCUMENT_BYTES + 1)
        doc = PlanDocuments(denial=denial)
        doc.plan_document_enc.save("huge.pdf", ContentFile(oversized), save=True)

        with patch.object(MLPlanDocHelper, "_extract_pages_with_terms") as extracted:
            result = asyncio.run(
                MLPlanDocHelper.extract_relevant_text(
                    denial.denial_id, ["medical necessity"]
                )
            )

        extracted.assert_not_called()
        self.assertEqual(result, "")

    def test_an_oversized_document_is_never_read_whole(self):
        """Skipping it after reading it costs the memory anyway.

        Rejecting on the length of what came back still allocates the whole
        file, and decryption allocates a second copy, on a worker other
        patients are generating on. So the limit has to reach the read.
        """
        from fighthealthinsurance.ml.ml_document_extraction import (
            read_and_decrypt_file,
        )

        class WatchedFile:
            """Stands in for the stored file and records how it is read."""

            def __init__(self, body):
                self.body = body
                self.asked_for = []

            def read(self, size=None):
                self.asked_for.append(size)
                return self.body if size is None else self.body[:size]

            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

        watched = WatchedFile(b"y" * 4096)
        field = mock.Mock()
        field.name = "huge.pdf"
        field.open.return_value = watched

        read_and_decrypt_file(field, 1024)

        self.assertEqual(
            watched.asked_for,
            [1025],
            "the read has to stop one byte past the ceiling: anything larger "
            "still allocates whatever the file happens to be",
        )

    def test_two_parses_never_overlap(self):
        """PyMuPDF is not safe to drive from several threads at once.

        Asserting the lock object exists proved nothing: removing the `with`
        around the parse left both assertions passing. This runs two parses
        from two threads and records when each is inside, so an unguarded
        parser overlaps and fails.
        """
        import threading
        import time

        from fighthealthinsurance.ml import ml_document_extraction

        inside = []
        overlapped = []
        real = ml_document_extraction.extract_text_from_plaintext_bytes

        def slow(data):
            inside.append(1)
            if len(inside) > 1:
                overlapped.append(1)
            time.sleep(0.05)
            inside.pop()
            return real(data)

        with mock.patch.object(
            ml_document_extraction, "extract_text_from_plaintext_bytes", slow
        ):
            threads = [
                threading.Thread(
                    target=ml_document_extraction.extract_text_from_bytes,
                    args=(b"plan text", "plan.txt"),
                )
                for _ in range(4)
            ]
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join()

        self.assertEqual(overlapped, [], "two parses ran at the same time")

    def test_the_parse_does_not_occupy_the_shared_executor(self):
        """A lock held on the default executor pins one of its threads.

        Every asyncio.to_thread call in the process shares that pool, so a
        few large documents could stall work with nothing to do with
        documents. Parsing has its own thread.

        Driven through the helper a plan document actually goes through, not
        through the wrapper directly: putting the production call back on the
        shared executor has to fail this, and calling the wrapper myself
        would not have noticed.
        """
        import asyncio

        from fighthealthinsurance.ml import ml_document_extraction

        seen = {}

        def note(data, filename):
            import threading

            seen["thread"] = threading.current_thread().name
            return "", {1: "medical necessity review"}

        async def run():
            with mock.patch.object(
                ml_document_extraction, "extract_text_from_bytes", note
            ):
                return await MLPlanDocHelper._extract_pages_with_terms(
                    b"plan bytes", "plan.pdf", ["medical necessity"]
                )

        pages = asyncio.run(run())

        self.assertTrue(pages, "the parse never happened, so nothing was proven")
        self.assertIn("fhi-doc-parse", seen["thread"], seen)

    def test_the_pdf_branch_itself_runs_under_the_lock(self):
        """The branch that crashes the worker is the one that must be held.

        Finding the `with` in the source proved nothing about where the
        parse happens: moving the PDF branch outside it left the string,
        and the overlap test above, both intact, because that test parses
        text. This asks the PDF parser whether the lock is held while it
        runs.
        """
        from fighthealthinsurance.ml import ml_document_extraction

        held = {}

        def note(data):
            held["locked"] = ml_document_extraction._PARSING.locked()
            return "", {}

        with mock.patch.object(
            ml_document_extraction, "extract_text_from_pdf_bytes", note
        ):
            ml_document_extraction.extract_text_from_bytes(b"%PDF-1.4", "plan.pdf")

        self.assertTrue(
            held.get("locked"), "the PDF parser ran with the lock free"
        )

    def test_a_cancelled_request_never_reaches_the_parser(self):
        """A queued work item keeps a decrypted document alive.

        Cancelling the future stops the parse, but the work item itself
        stays on the executor's queue until a worker dequeues it, and that
        item holds the decrypted bytes. A request that gave up waiting could
        therefore leave somebody's plan document in memory behind a parse
        still running. Waiting for the parser in the coroutine keeps the
        bytes in the request that owns them, so the queue is what this
        checks: nothing of a cancelled request is sitting in it.
        """
        import asyncio
        import contextlib
        import threading

        from fighthealthinsurance.ml import ml_document_extraction

        parsed: list[bytes] = []
        running = threading.Event()
        release = threading.Event()

        def blocking(data, filename):
            parsed.append(data)
            running.set()
            release.wait(10)
            return "", {}

        async def run():
            with mock.patch.object(
                ml_document_extraction, "extract_text_from_bytes", blocking
            ):
                first = asyncio.create_task(
                    ml_document_extraction.aextract_text_from_bytes(
                        b"the first document", "a.txt"
                    )
                )
                await asyncio.to_thread(running.wait, 10)
                second = asyncio.create_task(
                    ml_document_extraction.aextract_text_from_bytes(
                        b"the document nobody is waiting for", "b.txt"
                    )
                )
                await asyncio.sleep(0.05)
                second.cancel()
                with contextlib.suppress(asyncio.CancelledError):
                    await second
                # Reaching into the executor's queue on purpose: it is the
                # thing that holds the bytes, and the public surface cannot
                # show what is waiting in it.
                queued = ml_document_extraction._PARSER._work_queue.qsize()
                release.set()
                await first
                await asyncio.sleep(0.2)
                return queued

        queued = asyncio.run(run())

        self.assertEqual(
            queued,
            0,
            "a cancelled request left a decrypted document on the queue",
        )
        self.assertEqual(
            parsed,
            [b"the first document"],
            "a cancelled request's document was parsed anyway",
        )


class EveryAcceptedUploadIsReadTest(TransactionTestCase):
    """The upload form takes any file.

    Anything the extractor does not recognise is read as text, which is what
    this path did before it started decrypting. A format that needs its own
    parser, a saved web page above all, is worth doing properly and is not
    done here: building a document tree from a patient's file is its own
    memory question and belongs in its own change.
    """

    def _text_from(self, name, body):
        denial = Denial.objects.create(
            hashed_email=Denial.get_hashed_email(f"{name}@example.com"),
            denial_text="Denied.",
        )
        doc = PlanDocuments(denial=denial)
        doc.plan_document_enc.save(name, ContentFile(body), save=True)
        return asyncio.run(
            MLPlanDocHelper.extract_relevant_text(
                denial.denial_id,
                ["medical necessity", "necesidad", "autorizaci\u00f3n"],
            )
        )

    def test_a_binary_file_reads_as_nothing_useful_without_raising(self):
        """A scanned page is bytes, not words, and needs OCR this does not do.

        Read as text it simply matches no term. What it must not do is throw,
        because the caller would swallow that and the patient would get the
        same empty result with no idea why.
        """
        from fighthealthinsurance.ml.ml_document_extraction import (
            extract_text_from_bytes,
        )

        full, pages = extract_text_from_bytes(
            b"\xff\xd8\xff\xe0 not text at all", "scan.jpeg"
        )

        self.assertIsInstance(full, str)
        self.assertIsInstance(pages, dict)
        self.assertEqual(
            self._text_from("scan.jpeg", b"\xff\xd8\xff\xe0 not text at all"), ""
        )

    def test_an_unrecognised_extension_is_read_as_text(self):
        """A legacy upload named .text used to be read and stopped being.

        The path had a catch-all that opened anything non-PDF as text. An
        extension list replaced it, so a file it did not name contributed
        nothing at all, silently.
        """
        text = self._text_from(
            "benefits.text", b"Coverage requires medical necessity review."
        )

        self.assertIn("medical necessity", text.lower())

    def test_markdown_is_read_as_the_text_it_is(self):
        text = self._text_from(
            "plan.md", b"# Plan\n\nServices need **medical necessity** review.\n"
        )

        self.assertIn("medical necessity", text.lower())


def _make_docx_bytes() -> bytes:
    """A document whose deadline is in a table, as real plans write them."""
    import io as _io

    import docx

    document = docx.Document()
    document.add_paragraph("Plan overview for medical necessity review.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Appeal deadline"
    table.rows[0].cells[1].text = "60 days from receipt of this notice"
    buffer = _io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class WhatComesOutOfADocxTest(TransactionTestCase):
    """A .docx is an archive, and most of a plan's rules are in its tables."""

    def test_a_table_is_read(self):
        """``doc.paragraphs`` walks the top level only.

        A plan that states its appeal deadline in a table, which is how
        plans state them, lost exactly the provision the letter needs.
        """
        from fighthealthinsurance.ml.ml_document_extraction import (
            extract_text_from_docx_bytes,
        )

        full, pages = extract_text_from_docx_bytes(_make_docx_bytes())

        self.assertIn("medical necessity", full.lower())
        self.assertIn("Appeal deadline", full)
        self.assertIn("60 days from receipt", full)
        self.assertTrue(pages)

    def test_an_archive_that_unpacks_too_large_is_not_opened(self):
        """Zip turns a small upload into an arbitrarily large one.

        The stored-bytes ceiling bounds the archive, not its contents, and
        this is the first path that hands a patient's upload to a real
        document parser rather than decoding it as text.
        """
        from fighthealthinsurance.ml import ml_document_extraction

        data = _make_docx_bytes()

        with mock.patch.object(
            ml_document_extraction, "MAX_DOCX_UNPACKED_BYTES", 16
        ):
            full, pages = ml_document_extraction.extract_text_from_docx_bytes(data)

        self.assertEqual(full, "")
        self.assertEqual(pages, {})

    def test_an_ordinary_document_is_still_read(self):
        """The bound is not a wall in front of real documents."""
        from fighthealthinsurance.ml.ml_document_extraction import (
            extract_text_from_docx_bytes,
        )

        full, _pages = extract_text_from_docx_bytes(_make_docx_bytes())

        self.assertIn("Plan overview", full)


class ADamagedDocumentDoesNotBecomeGarbageTest(TransactionTestCase):
    """The text guess is for formats with no parser, not for failed parses."""

    def test_a_pdf_that_parses_to_nothing_is_not_decoded_as_text(self):
        """A scanned or damaged PDF must not contribute PDF syntax.

        Falling back on an empty result regardless of extension put the
        file's own internals into the summary whenever a search term
        happened to appear in them, which is the bug this branch exists to
        fix arriving by a different door.
        """
        broken = (
            b"%PDF-1.4\n% not a real pdf\n"
            b"/Title (medical necessity coverage policy)\n"
            b"endobj trailer"
        )

        pages = asyncio.run(
            MLPlanDocHelper._extract_pages_with_terms(
                broken, "plan.pdf", ["medical necessity"]
            )
        )

        self.assertEqual(pages, [])

    def test_a_format_with_no_parser_is_still_read_as_text(self):
        """The guess itself stays, for the uploads it was there for."""
        pages = asyncio.run(
            MLPlanDocHelper._extract_pages_with_terms(
                b"Coverage requires medical necessity review.",
                "benefits.text",
                ["medical necessity"],
            )
        )

        self.assertEqual(len(pages), 1)
        self.assertIn("medical necessity", pages[0].lower())
