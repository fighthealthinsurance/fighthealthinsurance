"""Shared plan documents are checked on the way in, and staff read them as text.

One upload takes up to SharePlanDocumentsForm.MAX_FILES files and
MAX_TOTAL_BYTES in all, which leaves room under the proxy's 80MB limit, so
nobody meets the proxy's error page. The page says so, its script says it
again before sending, and people with more are told to send another upload.

Every shared file gets a closer look than its first bytes: a .docx must be a
Word document without macros or locked parts, a .txt must be text, and a PDF
must open, without a password, in a separate process. The form carries a
signed note of when it was drawn, so a post that never loaded the page, or
came back faster than a person can choose files, is asked to try again before
any file is read. Each visitor, and everyone together, has a daily budget.

Staff read a shared file as text first, as plain text that runs nothing.
"""

import datetime
import io
import re
import time
import zipfile
from pathlib import Path
from unittest.mock import patch

from bs4 import BeautifulSoup
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from loguru import logger

from fighthealthinsurance import chat_forms, sentry_filters, shared_plan_documents
from fighthealthinsurance.chat_forms import SharePlanDocumentsForm
from fighthealthinsurance.models import PolicyDocument, SharedPlanDocumentsDay
from fighthealthinsurance.shared_plan_documents import (
    DAILY_LIMIT_MESSAGE,
    SESSION_LIMIT_MESSAGE,
    TRY_AGAIN,
    render_token,
)

User = get_user_model()

REPO = Path(__file__).resolve().parents[2]
TXT = b"Summary of Benefits and Coverage for the Crumpler Cooperative plan."


def txt(name="sbc.txt", content=TXT):
    return SimpleUploadedFile(name, content, content_type="text/plain")


def docx_bytes(text="Summary Plan Description for the Crumpler Cooperative"):
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def zip_bytes(entries):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return out.getvalue()


def zip_entries(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {entry.filename: archive.read(entry) for entry in archive.infolist()}


def docx_with_a_table():
    import docx

    document = docx.Document()
    document.add_paragraph("Summary Plan Description for the Crumpler Cooperative")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Deductible"
    table.cell(0, 1).text = "$500"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Out-of-pocket limit"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def docx_with_a_damaged_part(part_name):
    """A Word document with an extra part whose bytes don't match its CRC."""
    entries = zip_entries(docx_bytes())
    entries[part_name] = b"damaged part payload"
    data = bytearray(zip_bytes(entries))
    data[data.index(b"damaged part payload")] ^= 0xFF
    return bytes(data)


def with_encrypted_flag(data: bytes) -> bytes:
    """The same zip with every entry marked encrypted (general purpose bit 0)."""
    data = bytearray(data)
    for match in re.finditer(b"PK\x01\x02", bytes(data)):
        data[match.start() + 8] |= 0x1
    for match in re.finditer(b"PK\x03\x04", bytes(data)):
        data[match.start() + 6] |= 0x1
    return bytes(data)


def pdf_bytes(
    pages=1, text="Evidence of Coverage, Crumpler Health Plan", password=None
):
    import pymupdf

    document = pymupdf.open()
    for number in range(pages):
        page = document.new_page()
        page.insert_text((72, 72), f"{text} page {number + 1}")
    if password:
        data = document.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            owner_pw=password + "-owner",
            user_pw=password,
        )
    else:
        data = document.tobytes()
    document.close()
    return data


class _ShareTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.url = reverse("share_plan_documents")

    def tearDown(self):
        # Rows roll back with the test; their stored files would not.
        for doc in PolicyDocument.objects.all():
            if doc.document_enc:
                doc.document_enc.delete(save=False)

    def share(self, files, client=None, rendered=None, **fields):
        data = {
            "plan_documents": files,
            "agree": "on",
            "rendered": (
                render_token(at=time.time() - 60) if rendered is None else rendered
            ),
        }
        data.update(fields)
        return (client or self.client).post(self.url, data)

    def assertStored(self, response, count=1):
        self.assertRedirects(response, self.url + "?shared=1")
        self.assertEqual(PolicyDocument.objects.count(), count)

    def assertRefused(self, response, words, count=0):
        self.assertEqual(response.status_code, 200)
        self.assertEqual(PolicyDocument.objects.count(), count)
        self.assertContains(response, words)


class OneUploadFitsUnderTheProxyLimitTest(_ShareTestCase):
    def test_the_total_leaves_room_under_every_proxy_limit(self):
        limits = [
            (REPO / "conf" / "nginx.default", r"client_max_body_size\s+(\d+)M"),
            *[
                (REPO / "k8s" / name, r"proxy-body-size:\s*(\d+)m")
                for name in ("deploy.yaml", "deploy_staging.yaml", "deploy_dev.yaml")
            ],
        ]
        for path, pattern in limits:
            with self.subTest(path=path.name):
                megabytes = int(re.search(pattern, path.read_text()).group(1))
                # Room for the form's other fields and the multipart framing.
                self.assertLessEqual(
                    SharePlanDocumentsForm.MAX_TOTAL_BYTES + 10 * 1024 * 1024,
                    megabytes * 1024 * 1024,
                )

    def test_files_that_add_up_to_too_much_are_refused(self):
        with patch.object(SharePlanDocumentsForm, "MAX_TOTAL_BYTES", 100):
            response = self.share([txt(content=b"a" * 60), txt(content=b"b" * 60)])
        self.assertRefused(response, "add up to more than")

    def test_files_that_fit_are_stored(self):
        with patch.object(SharePlanDocumentsForm, "MAX_TOTAL_BYTES", 120):
            response = self.share([txt(content=b"a" * 60), txt(content=b"b" * 60)])
        self.assertStored(response, count=2)

    def test_the_count_stays_at_ten_files(self):
        self.assertEqual(SharePlanDocumentsForm.MAX_FILES, 10)


class ThePageSaysHowMuchFitsTest(_ShareTestCase):
    def setUp(self):
        super().setUp()
        self.soup = BeautifulSoup(self.client.get(self.url).content, "html.parser")

    def test_the_hint_says_how_many_files_and_how_much_fit(self):
        hint = self.soup.find(id="plan-documents-limits").get_text(" ", strip=True)
        self.assertIn("Up to 10 files and 60MB at a time", hint)
        self.assertIn("Have more? Send them in another upload.", hint)

    def test_the_script_gets_the_same_limits_and_words_as_the_server(self):
        form = self.soup.find("form", id="share-plan-documents-form")
        self.assertEqual(form["data-max-files"], str(SharePlanDocumentsForm.MAX_FILES))
        self.assertEqual(
            form["data-max-bytes"], str(SharePlanDocumentsForm.MAX_TOTAL_BYTES)
        )
        self.assertEqual(
            form["data-too-many"], SharePlanDocumentsForm.too_many_message()
        )
        self.assertEqual(form["data-too-big"], SharePlanDocumentsForm.too_big_message())
        # The script writes into the file field's own message, read out as it fills.
        self.assertIsNotNone(self.soup.find(id="id_plan_documents_error"))

    def test_the_thanks_offer_an_empty_form_for_more(self):
        soup = BeautifulSoup(
            self.client.get(self.url + "?shared=1").content, "html.parser"
        )
        thanks = soup.find(id="share-thanks")
        self.assertIn("Have more?", thanks.get_text(" ", strip=True))
        self.assertEqual(thanks.find("a")["href"], self.url)

    def test_each_page_carries_a_fresh_signed_render_time(self):
        rendered = self.soup.find("input", attrs={"name": "rendered"})
        self.assertEqual(rendered["type"], "hidden")
        self.assertTrue(rendered["value"])


class TheFormMustHaveBeenDrawnTest(_ShareTestCase):
    def test_a_post_without_the_render_time_is_asked_to_try_again(self):
        response = self.share([txt()], rendered="")
        self.assertRefused(response, TRY_AGAIN)

    def test_a_post_faster_than_a_person_is_asked_to_try_again(self):
        response = self.share([txt()], rendered=render_token())
        self.assertRefused(response, TRY_AGAIN)

    def test_a_page_left_open_for_days_is_asked_to_try_again(self):
        response = self.share(
            [txt()], rendered=render_token(at=time.time() - 3 * 86400)
        )
        self.assertRefused(response, TRY_AGAIN)

    def test_a_render_time_we_did_not_sign_is_asked_to_try_again(self):
        forged = render_token(at=time.time() - 60)[:-2] + "xx"
        response = self.share([txt()], rendered=forged)
        self.assertRefused(response, TRY_AGAIN)

    def test_asking_again_keeps_what_was_typed_and_reads_no_file(self):
        with patch.object(chat_forms, "validate_shared_document") as checked:
            response = self.share(
                [txt()], rendered="", insurer_name="Crumpler Mutual", plan_year="2025"
            )
        checked.assert_not_called()
        soup = BeautifulSoup(response.content, "html.parser")
        self.assertEqual(soup.find(id="id_insurer_name")["value"], "Crumpler Mutual")
        self.assertEqual(soup.find(id="id_plan_year")["value"], "2025")

    def test_the_hidden_field_is_checked_before_anything_else(self):
        with patch.object(chat_forms, "validate_shared_document") as checked:
            response = self.share([txt()], rendered="", website="https://spam.example")
        checked.assert_not_called()
        self.assertRedirects(response, self.url + "?shared=1")
        self.assertEqual(PolicyDocument.objects.count(), 0)


class EveryFileGetsACloserLookTest(_ShareTestCase):
    def test_a_word_document_is_stored(self):
        self.assertStored(self.share([SimpleUploadedFile("spd.docx", docx_bytes())]))

    def test_a_zip_that_is_not_a_word_document_is_refused(self):
        fake = zip_bytes({"readme.txt": b"not a document"})
        response = self.share([SimpleUploadedFile("spd.docx", fake)])
        self.assertRefused(response, "does not appear to be a valid DOCX")

    def test_a_word_document_with_macros_is_refused(self):
        with_macros = zip_bytes(
            {
                "[Content_Types].xml": b"<Types/>",
                "word/document.xml": b"<w:document/>",
                "word/vbaProject.bin": b"\x00macro",
            }
        )
        response = self.share([SimpleUploadedFile("spd.docx", with_macros)])
        self.assertRefused(response, "has macros")

    def test_a_word_document_with_locked_parts_is_refused(self):
        locked = with_encrypted_flag(docx_bytes())
        response = self.share([SimpleUploadedFile("spd.docx", locked)])
        self.assertRefused(response, "locked with a password")

    def test_a_password_protected_or_old_word_file_gets_its_own_message(self):
        ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
        response = self.share([SimpleUploadedFile("spd.docx", ole)])
        self.assertRefused(response, "is an old .doc file")

    def test_a_text_file_that_is_really_binary_is_refused(self):
        response = self.share([txt(content=b"MZ\x90\x00\x03\x00\x00\x00")])
        self.assertRefused(response, "doesn&#x27;t look like a text file")

    def test_a_pdf_is_stored(self):
        self.assertStored(self.share([SimpleUploadedFile("eoc.pdf", pdf_bytes())]))

    def test_a_pdf_locked_with_a_password_is_refused(self):
        locked = pdf_bytes(password="crumpler")
        response = self.share([SimpleUploadedFile("eoc.pdf", locked)])
        self.assertRefused(response, "This PDF is locked with a password.")

    def test_a_pdf_with_too_many_pages_is_refused(self):
        with patch.object(chat_forms, "MAX_SHARED_PDF_PAGES", 2):
            response = self.share([SimpleUploadedFile("eoc.pdf", pdf_bytes(pages=3))])
        self.assertRefused(response, "more than 2 pages")

    def test_a_pdf_the_strict_reader_trips_on_is_still_taken(self):
        # pypdf gives up on files the reader used for analysis can still open.
        self.assertStored(self.share([SimpleUploadedFile("eoc.pdf", b"%PDF-1.4 x")]))

    def test_a_pdf_that_takes_too_long_to_open_is_refused(self):
        with patch.object(chat_forms, "_PDF_CHECK_SECONDS", 0):
            response = self.share([SimpleUploadedFile("eoc.pdf", pdf_bytes())])
        self.assertRefused(response, "We couldn&#x27;t open this PDF.")

    def test_explain_policy_keeps_its_own_lighter_check(self):
        # The closer look is only for files we keep for research.
        fake = SimpleUploadedFile("policy.docx", b"PK\x03\x04 fake docx content")
        chat_forms.validate_policy_document(fake)
        with self.assertRaises(chat_forms.forms.ValidationError):
            chat_forms.validate_shared_document(fake)


class TheEmailIsHashedInTheFormTest(TestCase):
    def test_the_cleaned_form_holds_only_the_hash(self):
        form = SharePlanDocumentsForm(
            data={"agree": "on", "email": " Rebecca.Crumpler@Test-FHI.com "},
            files={"plan_documents": txt()},
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertNotIn("email", form.cleaned_data)
        self.assertEqual(len(form.cleaned_data["hashed_email"]), 128)
        self.assertNotIn("rebecca", repr(form.cleaned_data).lower())


class DailyBudgetsTest(_ShareTestCase):
    @override_settings(SHARED_PLAN_DOCUMENTS_SESSION_DAILY_FILES=2)
    def test_one_visitor_can_share_up_to_their_budget_in_a_day(self):
        self.assertStored(self.share([txt(), txt()]), count=2)
        self.assertRefused(self.share([txt()]), "share the rest tomorrow", count=2)

    @override_settings(SHARED_PLAN_DOCUMENTS_SESSION_DAILY_BYTES=100)
    def test_one_visitor_has_a_budget_in_bytes_too(self):
        self.assertStored(self.share([txt(content=b"a" * 80)]))
        self.assertRefused(
            self.share([txt(content=b"b" * 30)]), "share the rest tomorrow", count=1
        )

    @override_settings(SHARED_PLAN_DOCUMENTS_DAILY_FILES=1)
    def test_everyone_together_has_a_budget_whatever_the_cookies(self):
        self.assertStored(self.share([txt()], client=Client()))
        response = self.share([txt()], client=Client())
        self.assertRefused(response, "Please try again tomorrow.", count=1)
        self.assertNotEqual(SESSION_LIMIT_MESSAGE, DAILY_LIMIT_MESSAGE)

    @override_settings(SHARED_PLAN_DOCUMENTS_DAILY_FILES=0)
    def test_a_share_over_budget_is_refused_before_any_file_is_read(self):
        with patch.object(chat_forms, "validate_shared_document") as checked:
            response = self.share([txt()], insurer_name="Crumpler Mutual")
        checked.assert_not_called()
        self.assertRefused(response, "Please try again tomorrow.")
        self.assertContains(response, 'value="Crumpler Mutual"')

    @override_settings(SHARED_PLAN_DOCUMENTS_DAILY_FILES=1)
    def test_yesterdays_shares_do_not_count_today(self):
        self.share([txt()], client=Client())
        SharedPlanDocumentsDay.objects.update(
            day=timezone.now().date() - datetime.timedelta(days=1)
        )
        self.assertStored(self.share([txt()], client=Client()), count=2)

    @override_settings(SHARED_PLAN_DOCUMENTS_DAILY_FILES=1)
    def test_two_shares_at_once_cannot_both_take_the_last_room(self):
        # Both get past the first check, as two posts arriving together do.
        with patch.object(
            shared_plan_documents, "daily_limit_problem", return_value=None
        ):
            self.assertStored(self.share([txt()], client=Client()))
            response = self.share([txt()], client=Client())
        self.assertRefused(response, "Please try again tomorrow.", count=1)

    def test_a_share_that_fails_to_save_gives_its_room_back(self):
        with patch.object(
            PolicyDocument, "save", side_effect=RuntimeError("storage went away")
        ):
            self.share([txt(), txt()])
        self.assertEqual(shared_plan_documents.shared_today(), 0)


class StaffReadSharedFilesAsTextTest(_ShareTestCase):
    def read_as_text(self, doc, user=None):
        client = Client()
        if user is not None:
            client.force_login(user)
        return client.get(reverse("staff_policy_document_text", kwargs={"pk": doc.pk}))

    def staff(self):
        return User.objects.create_user(
            username="sarah_loguen_fraser", password="pw-for-tests", is_staff=True
        )

    def shared(self, upload):
        self.share([upload])
        return PolicyDocument.objects.get()

    def test_staff_read_a_text_file(self):
        doc = self.shared(txt())
        response = self.read_as_text(doc, self.staff())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/plain; charset=utf-8")
        self.assertIn(TXT.decode(), response.content.decode())

    def test_staff_read_a_pdf_page_by_page(self):
        doc = self.shared(SimpleUploadedFile("eoc.pdf", pdf_bytes(pages=2)))
        text = self.read_as_text(doc, self.staff()).content.decode()
        self.assertIn("[Page 2]", text)
        self.assertIn("Crumpler Health Plan page 2", text)

    def test_staff_read_a_word_document(self):
        doc = self.shared(SimpleUploadedFile("spd.docx", docx_bytes()))
        text = self.read_as_text(doc, self.staff()).content.decode()
        self.assertIn("Summary Plan Description for the Crumpler Cooperative", text)

    def test_staff_read_a_word_table_in_one_pass_over_its_text(self):
        doc = self.shared(SimpleUploadedFile("spd.docx", docx_with_a_table()))
        with patch(
            "docx.Document", side_effect=AssertionError("parsed by python-docx")
        ):
            text = self.read_as_text(doc, self.staff()).content.decode()
        self.assertIn("Deductible\n$500\nOut-of-pocket limit", text)

    def test_reading_a_word_file_logs_nothing_from_inside_it(self):
        damaged = docx_with_a_damaged_part("rebecca.crumpler@test-fhi.com.xml")
        doc = self.shared(SimpleUploadedFile("spd.docx", damaged))
        logged = []
        sink = logger.add(lambda message: logged.append(str(message)), level="DEBUG")
        try:
            text = self.read_as_text(doc, self.staff()).content.decode()
        finally:
            logger.remove(sink)
        self.assertNotIn("rebecca.crumpler", "".join(logged).lower())
        self.assertIn("Summary Plan Description for the Crumpler Cooperative", text)

    def test_markup_in_a_file_is_shown_as_text_that_runs_nothing(self):
        doc = self.shared(txt(content=b"<script>alert('Crumpler')</script>"))
        response = self.read_as_text(doc, self.staff())
        self.assertEqual(response["Content-Type"], "text/plain; charset=utf-8")
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        self.assertEqual(
            response["Content-Security-Policy"], "default-src 'none'; sandbox"
        )
        self.assertIn("no-store", response["Cache-Control"])

    def test_people_who_are_not_staff_cannot_read_it(self):
        doc = self.shared(txt())
        member = User.objects.create_user(
            username="mary_eliza_mahoney", password="pw-for-tests", is_staff=False
        )
        for response in (self.read_as_text(doc, member), self.read_as_text(doc)):
            with self.subTest(status=response.status_code):
                self.assertEqual(response.status_code, 302)
                self.assertNotIn(b"Crumpler", response.content)

    def test_the_raw_download_runs_nothing_either(self):
        doc = self.shared(txt())
        client = Client()
        client.force_login(self.staff())
        response = client.get(
            reverse("staff_policy_document_download", kwargs={"pk": doc.pk})
        )
        self.assertEqual(
            response["Content-Security-Policy"], "default-src 'none'; sandbox"
        )
        self.assertIn("attachment", response["Content-Disposition"])

    def test_the_admin_offers_the_text_before_the_original(self):
        doc = self.shared(txt())
        admin = User.objects.create_superuser(
            username="susan_la_flesche_picotte", password="pw-for-tests"
        )
        client = Client()
        client.force_login(admin)
        page = client.get(
            reverse("admin:fighthealthinsurance_policydocument_change", args=[doc.pk])
        ).content.decode()
        text_link = reverse("staff_policy_document_text", kwargs={"pk": doc.pk})
        download = reverse("staff_policy_document_download", kwargs={"pk": doc.pk})
        self.assertLess(page.index(text_link), page.index(download))
        self.assertNotIn('name="document_enc"', page)


class ErrorReportsFromTheSharePageCarryNoFormDataTest(TestCase):
    def event(self, url):
        return {
            "request": {"url": url, "data": {"email": "rebecca@test-fhi.com"}},
            "exception": {
                "values": [
                    {
                        "type": "RuntimeError",
                        "value": "storage went away",
                        "stacktrace": {
                            "frames": [{"function": "f", "vars": {"email": "x"}}]
                        },
                    }
                ]
            },
        }

    def frames(self, event):
        return event["exception"]["values"][0]["stacktrace"]["frames"]

    def test_an_error_on_the_share_page_loses_the_form_and_the_variables(self):
        event = sentry_filters.before_send_filter(
            self.event("https://www.fighthealthinsurance.com/share-plan-documents"), {}
        )
        self.assertNotIn("data", event["request"])
        self.assertNotIn("vars", self.frames(event)[0])
        self.assertEqual(self.frames(event)[0]["function"], "f")

    def test_an_error_elsewhere_keeps_them(self):
        event = sentry_filters.before_send_filter(
            self.event("https://www.fighthealthinsurance.com/share-plan-documents-x"),
            {},
        )
        self.assertIn("data", event["request"])
        self.assertIn("vars", self.frames(event)[0])
